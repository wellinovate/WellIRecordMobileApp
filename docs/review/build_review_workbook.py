import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

d=json.load(open('data.json')); imm=d['imm']; ms=d['ms']
F='Arial'
wb=Workbook()
hdr=PatternFill('solid',fgColor='082F67'); fill_in=PatternFill('solid',fgColor='FFF9DB'); thin=Side(style='thin',color='CCCCCC')
box=Border(left=thin,right=thin,top=thin,bottom=thin)
def head(ws,row,cols,widths):
    for i,(c,w) in enumerate(zip(cols,widths),1):
        x=ws.cell(row=row,column=i,value=c); x.font=Font(name=F,bold=True,color='FFFFFF'); x.fill=hdr; x.alignment=Alignment(wrap_text=True,vertical='center'); x.border=box
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes=ws.cell(row=row+1,column=3)

# --- Read me
ws=wb.active; ws.title='Read me'
rows=[
 ('WelliRecord clinical review: child immunization schedule and developmental milestones',True),
 ('',False),
 ('What this is',True),
 ('The app shows parents a vaccine schedule and a developmental checklist. Both are drafts compiled from public sources. They are marked "pending clinical review" in the app until a clinician signs them off.',False),
 ('What we ask of you',True),
 ('1. Check each row on the two data tabs against the official Nigerian source you trust (the current NPHCDA schedule for vaccines; your clinical judgment for milestones).',False),
 ('2. In the yellow columns, pick Correct, Change or Remove. If Change, write the correct value and the source. Add rows for anything missing under "Missing items".',False),
 ('3. Fill in the Sign-off tab. We record your name and date in the app. We do not publish the schedule as reviewed until you have signed.',False),
 ('How to read the vaccine columns',True),
 ('"Due (days after birth)" is when the dose first counts as due. "Grace (days)" is how long after that before the app calls it overdue. Programme doses (track = programme) are never flagged overdue; they are only recorded.',False),
 ('Example of a completed row (not real data)',True),
 ('Row: BCG-1 | Correct? = Change | Corrected value = "Grace 8 weeks" | Source = "NPHCDA RI schedule, 2023 ed." | Comment = "Catch-up allowed to 1 year".',False),
 ('Limits',True),
 ('The milestone list is from the US CDC "Learn the Signs. Act Early" checklists. We need to know whether each item is appropriate for Nigerian children and language, and whether wording is clear in plain English. The app never gives a result or diagnosis from it.',False),
]
for r,(t,b) in enumerate(rows,1):
    c=ws.cell(row=r,column=1,value=t); c.font=Font(name=F,bold=b,size=13 if r==1 else 11); c.alignment=Alignment(wrap_text=True,vertical='top')
ws.column_dimensions['A'].width=120

def days_text(n):
    if n==0: return 'At birth'
    if n<150: return f'{round(n/7)} weeks'
    if n<730: return f'{round(n/30)} months'
    return f'{round(n/365)} years'

# --- Immunization
ws=wb.create_sheet('Vaccines')
ws['A1']=f"Immunization schedule, version {imm['version']}, status: {imm['reviewStatus']}"; ws['A1'].font=Font(name=F,bold=True)
ws['A2']='Sources used: '+'; '.join(imm['sources']); ws['A2'].font=Font(name=F,italic=True); ws['A2'].alignment=Alignment(wrap_text=True)
ws.merge_cells('A2:N2'); ws.row_dimensions[2].height=48
cols=['Code','Vaccine','Protects against','Dose #','Due (days after birth)','Due (plain)','Grace (days)','Track','Route','Note in app','Correct?','Corrected value','Source cited','Comment']
head(ws,4,cols,[11,28,34,7,12,12,10,11,14,44,12,28,32,36])
r=5
for x in imm['doses']:
    vals=[x['code'],x['vaccine'],x['antigen'],x['doseNumber'],x['dueDays'],days_text(x['dueDays']),x['graceDays'],x['track'],x.get('route',''),x.get('note',''),None,None,None,None]
    for i,v in enumerate(vals,1):
        c=ws.cell(row=r,column=i,value=v); c.font=Font(name=F); c.border=box; c.alignment=Alignment(wrap_text=True,vertical='top')
        if i>=11: c.fill=fill_in
    r+=1
last=r-1
dv=DataValidation(type='list',formula1='"Correct,Change,Remove"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'K5:K{last}')
r+=1
ws.cell(row=r,column=1,value='Missing items (doses the schedule should have and does not)').font=Font(name=F,bold=True)
for k in range(1,6):
    for i in range(1,15):
        c=ws.cell(row=r+k,column=i); c.border=box; c.fill=fill_in; c.font=Font(name=F)
# summary with formulas
sr=r+8
ws.cell(row=sr,column=1,value='Review progress').font=Font(name=F,bold=True)
for j,(lab,f) in enumerate([('Rows',f'=COUNTA(A5:A{last})'),('Marked Correct',f'=COUNTIF(K5:K{last},"Correct")'),('Marked Change',f'=COUNTIF(K5:K{last},"Change")'),('Marked Remove',f'=COUNTIF(K5:K{last},"Remove")'),('Still to review',f'=B{sr+1}-B{sr+2}-B{sr+3}-B{sr+4}')],1):
    ws.cell(row=sr+j,column=1,value=lab).font=Font(name=F); ws.cell(row=sr+j,column=2,value=f).font=Font(name=F)

# --- Milestones
ws=wb.create_sheet('Milestones')
ws['A1']=f"Milestone checklist, version {ms['version']}, status: {ms['reviewStatus']}"; ws['A1'].font=Font(name=F,bold=True)
ws['A2']='Source: '+ms['source']; ws['A2'].font=Font(name=F,italic=True)
cols=['Code','Age (months)','Area','Milestone (as shown to parents)','Appropriate for Nigeria?','Wording fix','Comment']
head(ws,4,cols,[10,10,24,66,16,44,36])
r=5
for x in ms['milestones']:
    vals=[x['code'],x['ageMonths'],ms['domains'][x['domain']],x['text'],None,None,None]
    for i,v in enumerate(vals,1):
        c=ws.cell(row=r,column=i,value=v); c.font=Font(name=F); c.border=box; c.alignment=Alignment(wrap_text=True,vertical='top')
        if i>=5: c.fill=fill_in
    r+=1
lastm=r-1
dv2=DataValidation(type='list',formula1='"Yes,Change,Remove"',allow_blank=True); ws.add_data_validation(dv2); dv2.add(f'E5:E{lastm}')
sr=lastm+2
ws.cell(row=sr,column=1,value='Review progress').font=Font(name=F,bold=True)
for j,(lab,f) in enumerate([('Rows',f'=COUNTA(A5:A{lastm})'),('Yes',f'=COUNTIF(E5:E{lastm},"Yes")'),('Change',f'=COUNTIF(E5:E{lastm},"Change")'),('Remove',f'=COUNTIF(E5:E{lastm},"Remove")'),('Still to review',f'=B{sr+1}-B{sr+2}-B{sr+3}-B{sr+4}')],1):
    ws.cell(row=sr+j,column=1,value=lab).font=Font(name=F); ws.cell(row=sr+j,column=2,value=f).font=Font(name=F)

# --- Sign-off
ws=wb.create_sheet('Sign-off')
ws.column_dimensions['A'].width=34; ws.column_dimensions['B'].width=60
items=['Reviewer full name','Professional title','Registration body and number (MDCN, etc.)','Organisation','Email','Date of review','Which tabs did you review? (Vaccines / Milestones / both)','Schedule version reviewed','Declaration: I have checked this against the sources named and the changes above are my professional advice.','Signature (typed name is fine)']
for i,t in enumerate(items,1):
    a=ws.cell(row=i,column=1,value=t); a.font=Font(name=F,bold=True); a.alignment=Alignment(wrap_text=True,vertical='top')
    b=ws.cell(row=i,column=2); b.fill=fill_in; b.border=box; b.font=Font(name=F)
ws['B8']=f"vaccines {imm['version']}; milestones {ms['version']}"
wb.save('WelliRecord_clinical_review.xlsx')
print(last-4,'vaccine rows',lastm-4,'milestone rows')
